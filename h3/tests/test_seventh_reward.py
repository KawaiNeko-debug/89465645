import importlib
import json
import threading
from concurrent.futures import ThreadPoolExecutor
from copy import deepcopy
from unittest.mock import Mock, patch

import pytest
from openpyxl import load_workbook

from h3.seventh_reward import (
    MonthlyQuota, StateConflict, BEANS_PATH, COUPON_PATH, LIST_PATH, RULE_PATH,
    SHELL_COUPON, applies, choose_reward, empty_reward, merge_reward,
    monthly_summary, safe_reward,
)
from h3.retry_components import retry_components
from h3.merge_results import pick_result
from h3.mixed_results import applicable_components, stamp_result, merge_results
from h3.report import build_message, write_xlsx


class MemoryStore:
    def __init__(self):
        self.files = {}
        self.version = 0
        self.lock = threading.RLock()

    def ensure_branch(self):
        pass

    def read(self, path):
        with self.lock:
            payload, sha = self.files.get(path, (None, None))
            return deepcopy(payload), sha

    def write(self, path, payload, sha=None):
        with self.lock:
            _, current = self.files.get(path, (None, None))
            if current != sha:
                raise StateConflict('changed')
            self.version += 1
            self.files[path] = deepcopy(payload), str(self.version)
            return str(self.version)


@pytest.fixture
def frozen_clock(monkeypatch):
    from datetime import datetime
    class FrozenDateTime(datetime):
        @classmethod
        def now(cls, tz=None):
            return cls(2026, 10, 14, 12, 0, tzinfo=tz)
    monkeypatch.setattr('h3.seventh_reward.datetime', FrozenDateTime)


def seed(quota, count):
    entries = {}
    for i in range(1, count + 1):
        entries[quota.key('wudi1', i, '2026-10-01')] = {
            'source_group': 'wudi1', 'account_index': i, 'month': quota.month,
            'claim_date': '2026-10-01', 'state': 'completed', 'quantity': 1,
            'reservation': f'{i:032x}',
        }
    quota.store.write(quota.path, {'month': quota.month, 'entries': entries})


def current(**overrides):
    return dict(signModeCode='sign002', day=7, haveSignIn=True,
                haveReceive=False, version=5, **overrides)


def fake_request(coupon_response=None, *, can_exchange=True, options=None, rule_scene=3):
    coupon_response = {'code': 200, 'data': None} if coupon_response is None else coupon_response
    options = ([{'id': 'dynamic-option', 'goodsTitle': 'PCB加急减免券',
                 'isCanExchange': can_exchange, 'terminal': 'WEB', 'rewardType': 'coupon'}]
               if options is None else options)
    return Mock(side_effect=lambda path, params: {
        RULE_PATH: {'code': 200, 'data': {'signInRuleConfigs': [{'day': 7, 'scene': rule_scene}]}},
        LIST_PATH: {'code': 200, 'data': options},
        COUPON_PATH: coupon_response,
        BEANS_PATH: {'success': True, 'data': 10},
    }[path])


@pytest.mark.parametrize('count', [0, 1, 29, 30, 31])
def test_actual_monthly_threshold(count, frozen_clock):
    quota = MonthlyQuota('2026-10', MemoryStore())
    seed(quota, count)
    request = fake_request()
    row = choose_reward('wudi2', 1, '2026-10-14', current(), request, quota)
    assert row['choice'] == ('coupon' if count < 30 else 'beans')
    assert row['complete']
    assert monthly_summary('2026-10', quota.store)['acquired'] == count + (count < 30)
    paths = [call.args[0] for call in request.call_args_list]
    assert (COUPON_PATH in paths) == (count < 30)
    if count < 30:
        claim = next(call for call in request.call_args_list if call.args[0] == COUPON_PATH)
        assert claim.args[1] == {'couponId': 'dynamic-option', 'terminal': 'WEB'}


def test_concurrent_last_slot_is_reserved_once(monkeypatch):
    monkeypatch.setattr('h3.seventh_reward.time.sleep', lambda _: None)
    quota = MonthlyQuota('2026-10', MemoryStore())
    seed(quota, 29)
    with ThreadPoolExecutor(max_workers=20) as pool:
        values = list(pool.map(lambda index: quota.reserve('wudi2', index, '2026-10-14'), range(1, 21)))
    assert sum(state == 'reserved' for state, _ in values) == 1
    assert MonthlyQuota.total(quota.read()[0]) == 29
    winner = next(i for i, (state, _) in enumerate(values, 1) if state == 'reserved')
    quota.update('wudi2', winner, '2026-10-14', values[winner - 1][1], 'completed')
    assert monthly_summary('2026-10', quota.store)['acquired'] == 30


@pytest.mark.parametrize('group', ['old1', 'ld1', 'yyy1', 'll1', 'zh1', 'test', 'gift_test', 'wudi21'])
def test_only_wudi_strategy(group, frozen_clock):
    request = Mock()
    assert not applies(group)
    assert not choose_reward(group, 1, '2026-10-14', current(), request)['required']
    request.assert_not_called()
    assert 'seventh_reward' not in applicable_components({'source_group': group}, '2026-10-14')


@pytest.mark.parametrize('message', ['您已领取过优惠券，在限制周期内无法再次参与', '奖励不可选', '重复领取金豆'])
def test_business_terminals_release_quota(message, frozen_clock):
    quota = MonthlyQuota('2026-10', MemoryStore())
    row = choose_reward('wudi1', 1, '2026-10-14', current(),
                        fake_request({'success': False, 'message': message}), quota)
    assert row['complete'] and row['period_limited']
    assert MonthlyQuota.total(quota.read()[0]) == 0
    assert quota.reserve('wudi2', 1, '2026-10-14')[0] == 'reserved'


def test_coupon_specific_cycle_limit_selects_beans_without_counting(frozen_clock):
    quota = MonthlyQuota('2026-10', MemoryStore())
    request = fake_request({'success': False, 'message': '优惠券兑换已上限'})
    row = choose_reward('wudi1', 1, '2026-10-14', current(), request, quota)
    assert row['complete'] and row['choice'] == 'beans' and row['period_limited']
    assert row['fallback'] and row['coupon_count'] == 0
    assert [call.args[0] for call in request.call_args_list][-1] == BEANS_PATH
    assert MonthlyQuota.total(quota.read()[0]) == 0


@pytest.mark.parametrize('options', [[], [{'id': 'old', 'goodsTitle': '加急减免券', 'isCanExchange': False}],
                                   [{'id': 'expired', 'goodsTitle': '加急减免券', 'endDate': '2026-10-13'}],
                                   [{'id': 'future', 'goodsTitle': '加急减免券', 'startDate': '2026-10-15'}],
                                   [{'id': 'bean', 'goodsTitle': '加急减免券', 'rewardType': 'bean'}]])
def test_unavailable_coupon_falls_back_to_beans(options, frozen_clock):
    request = fake_request(options=options)
    row = choose_reward('wudi1', 1, '2026-10-14', current(), request, MonthlyQuota('2026-10', MemoryStore()))
    assert row['choice'] == 'beans' and row['fallback'] and row['complete']
    assert COUPON_PATH not in [call.args[0] for call in request.call_args_list]


def test_explicit_failure_retries_only_reward_and_releases_slot(frozen_clock):
    quota = MonthlyQuota('2026-10', MemoryStore())
    row = choose_reward('wudi1', 1, '2026-10-14', current(),
                        fake_request({'success': False, 'message': '服务器繁忙'}), quota)
    assert not row['complete'] and row['interface_failed']
    completed = complete_record(row)
    assert retry_components(completed) == ['seventh_reward']
    better = choose_reward('wudi1', 1, '2026-10-14', current(), fake_request(), quota, previous=row)
    assert better['complete'] and better['coupon_count'] == 1
    assert MonthlyQuota.total(quota.read()[0]) == 1


def test_lost_response_never_resubmits_or_counts(frozen_clock):
    quota = MonthlyQuota('2026-10', MemoryStore())
    request = fake_request()
    original = request.side_effect
    request.side_effect = lambda path, params: None if path == COUPON_PATH else original(path, params)
    row = choose_reward('wudi1', 1, '2026-10-14', current(), request, quota)
    assert row['complete'] and row['status'] == '领取结果未知'
    assert monthly_summary('2026-10', quota.store) == {
        'month': '2026-10', 'acquired': 0, 'remaining': 30, 'unknown': 1}
    assert retry_components(complete_record(row)) == []
    second_request = fake_request()
    # Even without the result artifact, the persistent reservation prevents duplicates.
    other = choose_reward('wudi1', 1, '2026-10-15', current(), second_request, quota)
    assert other['status'] == '领取结果未知'
    assert COUPON_PATH not in [call.args[0] for call in second_request.call_args_list]


def test_state_failure_after_confirmation_only_retries_checkpoint(frozen_clock):
    quota = MonthlyQuota('2026-10', MemoryStore())
    update = quota.update
    def fail_confirm(*args):
        if args[-1] == 'completed':
            raise RuntimeError('synthetic checkpoint failure')
        return update(*args)
    quota.update = fail_confirm
    row = choose_reward('wudi1', 1, '2026-10-14', current(), fake_request(), quota)
    assert row['coupon_count'] == 1 and not row['complete']
    quota.update = update
    request = Mock()
    cfg = current()
    cfg['haveReceive'] = True
    repaired = choose_reward('wudi1', 1, '2026-10-14', cfg, request, quota, previous=row)
    assert repaired['complete']
    request.assert_not_called()
    assert MonthlyQuota.total(quota.read()[0]) == 1


def test_confirmed_claim_proof_survives_config_read_outage(frozen_clock):
    previous = empty_reward()
    previous.update(required=True, complete=False, coupon_count=1,
                    month='2026-10', claim_date='2026-10-14', status='状态保存失败')
    result = choose_reward('wudi1', 1, '2026-10-14', None, Mock(), previous=previous)
    assert result['coupon_count'] == 1 and result['claim_date'] == '2026-10-14'
    assert not result['complete']


def test_reservation_can_transfer_only_before_claim_submission():
    quota = MonthlyQuota('2026-10', MemoryStore())
    _, old = quota.reserve('wudi1', 1, '2026-10-14')
    state, new = quota.reserve('wudi1', 1, '2026-10-14')
    assert state == 'reserved' and new != old
    with pytest.raises(StateConflict):
        quota.update('wudi1', 1, '2026-10-14', old, 'sent')
    quota.update('wudi1', 1, '2026-10-14', new, 'sent')
    assert quota.reserve('wudi1', 1, '2026-10-15')[0] == 'unknown'


def test_state_outage_never_submits_a_claim(frozen_clock):
    store = MemoryStore()
    store.ensure_branch = Mock(side_effect=RuntimeError('unavailable'))
    request = fake_request()
    row = choose_reward('wudi1', 1, '2026-10-14', current(), request, MonthlyQuota('2026-10', store))
    assert not row['complete'] and row['status'] == '状态保存失败'
    assert COUPON_PATH not in [call.args[0] for call in request.call_args_list]


@pytest.mark.parametrize('path', [RULE_PATH, LIST_PATH])
def test_read_failure_preserves_existing_components(path, frozen_clock):
    request = fake_request()
    original = request.side_effect
    request.side_effect = lambda requested, params: None if requested == path else original(requested, params)
    row = choose_reward('wudi1', 1, '2026-10-14', current(), request, MonthlyQuota('2026-10', MemoryStore()))
    assert retry_components(complete_record(row)) == ['seventh_reward']
    assert COUPON_PATH not in [call.args[0] for call in request.call_args_list]


def test_already_claimed_and_non_seven_config_do_not_submit(frozen_clock):
    request = Mock()
    cfg = current()
    cfg['haveReceive'] = True
    row = choose_reward('wudi1', 1, '2026-10-14', cfg, request)
    assert row['complete'] and row['status'] == '奖励已领取' and row['coupon_count'] == 0
    cfg['signModeCode'] = 'sign001'
    assert not choose_reward('wudi1', 1, '2026-10-14', cfg, request)['required']
    request.assert_not_called()


def test_other_group_and_skip_sign_ignore_stale_reward_failure():
    reward = empty_reward()
    reward.update(required=True, complete=False, status='接口失败')
    row = complete_record(reward)
    row['group_code'] = 'ld1'
    assert retry_components(row) == []
    row.update(group_code='wudi1', sign_skipped=True)
    assert retry_components(row) == []


def test_no_coupon_rule_uses_beans(frozen_clock):
    request = fake_request(rule_scene=1)
    row = choose_reward('wudi1', 1, '2026-10-14', current(), request, MonthlyQuota('2026-10', MemoryStore()))
    assert row['choice'] == 'beans' and row['complete']
    assert LIST_PATH not in [call.args[0] for call in request.call_args_list]


def test_month_rollover_and_confirmed_idempotency(frozen_clock):
    store = MemoryStore()
    october = MonthlyQuota('2026-10', store)
    seed(october, 30)
    november = MonthlyQuota('2026-11', store)
    assert november.reserve('wudi2', 1, '2026-11-01')[0] == 'reserved'
    assert monthly_summary('2026-11', store)['acquired'] == 0
    previous = empty_reward()
    previous.update(required=True, complete=True, coupon_count=1, status='领取成功')
    request = Mock()
    assert choose_reward('wudi1', 1, '2026-10-14', current(), request, october, previous) == previous
    request.assert_not_called()


def complete_record(reward=None):
    return {'group_code': 'wudi1', 'source_group': 'wudi1', 'account_index': 1,
            'sign_success': True, 'sign_status': '签到成功', 'points_fetch_success': True,
            'activity_fetch_success': True, 'data_fetch_completed': True,
            'account_data': {'invoice_fetch_success': True, 'pcb_order_fetch_success': True,
                             'coupon_fetch_success': True, 'balance_fetch_success': True},
            'seventh_reward': reward or empty_reward(), 'task_start_date': '2026-10-14'}


def test_successful_reward_survives_poor_retry_and_artifact_merge(tmp_path):
    acquired = empty_reward()
    acquired.update(required=True, coupon_count=1, status='领取成功', month='2026-10')
    failed = empty_reward()
    failed.update(required=True, complete=False, status='接口失败', interface_failed=True)
    assert merge_reward(failed, acquired) == acquired
    assert pick_result(complete_record(acquired), complete_record(failed))['seventh_reward'] == acquired
    result_dir = tmp_path / 'first'
    result_dir.mkdir()
    (result_dir / 'result.json').write_text(json.dumps({'results': [complete_record(acquired)]}), encoding='utf-8')
    output = tmp_path / 'merged.json'
    merge_results(tmp_path, output, [{'source_group': 'wudi1', 'account_index': 1}], '2026-10-14')
    assert json.loads(output.read_text(encoding='utf-8'))['results'][0]['seventh_reward']['coupon_count'] == 1


def test_projection_never_serializes_server_response_or_credentials():
    value = {'status': 'server body containing secret', 'cookie': 'synthetic-cookie',
             'couponId': 'synthetic-id', 'password': 'synthetic-password', 'username': 'synthetic-user'}
    encoded = json.dumps(safe_reward(value))
    assert all(word not in encoded for word in ('synthetic-', 'server body', 'cookie', 'password', 'username', 'couponId'))


def test_priority_sheets_validity_colors_and_row_sort(tmp_path, monkeypatch):
    monkeypatch.setenv('CONFIDENTIAL_PASSWORDS', 'synthetic-hidden')
    records = []
    for index, (invoice, password) in enumerate([('无', 'synthetic-hidden'), ('有', 'synthetic-hidden'),
                                               ('无', 'synthetic-visible'), ('有', 'synthetic-visible')], 1):
        record = complete_record()
        record.update(account_index=index, username=f'fixture-{index}', password=password)
        record['account_data'].update(invoice_profile_status=invoice, coupons={'unused': [
            {'name': 'PCB+SMT优惠券', 'coupon_extend_state': 'PCB+SMT', 'expires_at': '2099-01-01'},
            {'name': 'PCB加急减免券', 'expires_at': '2099-01-01'},
            {'name': SHELL_COUPON, 'expires_at': '2099-01-01'},
            {'name': '普通1折券'}, {'name': '普通券'},
            {'name': '过期加急减免券', 'expires_at': '2000-01-01'},
            {'name': '未来加急减免券', 'valid_from': '2099-01-01'},
        ], 'used': [{'name': '已用加急减免券'}], 'expired': []})
        records.append(record)
    path = tmp_path / 'report.xlsx'
    write_xlsx(str(path), records)
    workbook = load_workbook(path)
    assert workbook.sheetnames[:4] == ['签到汇总', 'PCB+SMT券', '加急减免券', SHELL_COUPON]
    assert not any('过期' in name or '未来' in name for name in workbook.sheetnames)
    target = workbook['加急减免券']
    assert [target.cell(row, 2).value for row in range(2, 6)] == ['fixture-4', 'fixture-2', 'fixture-3', 'fixture-1']
    assert target.cell(2, 4).font.color.rgb == '00008000'
    assert target.cell(4, 4).font.color.rgb == '00C00000'
    assert target.cell(3, 3).value == '保密'
    assert target.max_row == 5
    empty = tmp_path / 'empty.xlsx'
    write_xlsx(str(empty), [complete_record()])
    assert load_workbook(empty).sheetnames == ['签到汇总']


def test_wudi_telegram_counts_without_password(monkeypatch):
    monkeypatch.setenv('SUMMARY_CATEGORY', '无敌全干组')
    reward = empty_reward()
    reward.update(required=True, choice='coupon', coupon_count=1, status='领取成功')
    row = complete_record(reward)
    row.update(username='fixture-account', password='synthetic-password')
    message, summary = build_message([row], {'seventh_reward_monthly': {
        'month': '2026-10', 'acquired': 29, 'remaining': 1, 'unknown': 0}}, 1)
    assert '本月加急减免券已领取: 29' in message
    assert '距离30张还差: 1' in message
    assert '选择加急减免券账号数: 1' in message
    assert 'synthetic-password' not in message
    monkeypatch.setenv('SUMMARY_CATEGORY', '老号全干组')
    row['group_code'] = 'old1'
    other, _ = build_message([row], {}, 1)
    assert '本月加急减免券' not in other


def script_module(monkeypatch):
    for key, value in {
        'BASE_URL': 'https://example.invalid', 'PASSPORT_URL': 'https://example.invalid',
        'REFERER': 'https://example.invalid', 'SLIDER_ID': 'slider', 'WRAPPER_ID': 'wrapper',
        'HEADER_CLIENT_TYPE': 'client', 'HEADER_ACCESS_TOKEN': 'access', 'TOKEN_KEY': 'token',
        'GROUP_CODE': 'wudi1',
    }.items():
        monkeypatch.setenv(key, value)
    return importlib.import_module('h3.script')


def test_sign_success_is_independent_of_coupon_failure(monkeypatch):
    script = script_module(monkeypatch)
    client = script.ApiClient('synthetic-access', '', 1, Mock())
    client.get_points = Mock(return_value=100)
    cfg = current()
    cfg.update(haveSignIn=False, haveReceive=True)
    client.get_sign_config = Mock(return_value=(False, 7, True, {'data': cfg}))
    def sign():
        client.sign_status = '签到成功'
        client.sign_completed_at = '2026-10-14 12:00:00'
        return True
    client.sign_in = Mock(side_effect=sign)
    def reward(*args):
        client.seventh_reward.update(required=True, complete=False, status='接口失败', interface_failed=True)
    client.execute_seventh_reward = Mock(side_effect=reward)
    assert client.execute_wudi_sign_process()
    assert retry_components(complete_record(client.seventh_reward)) == ['seventh_reward']
    client.sign_in.assert_called_once()


def test_official_claim_request_suppresses_all_sensitive_output(monkeypatch, capsys):
    script = script_module(monkeypatch)
    client = script.ApiClient('synthetic-access', '', 1, Mock())
    with patch.object(script.requests, 'get', side_effect=RuntimeError('synthetic-cookie synthetic-password')):
        assert client.reward_request(COUPON_PATH, {'couponId': 'synthetic-id'}) is None
    assert capsys.readouterr().out == ''
